"""CloudCare synthetic raw operational data. Python 3.10+, numpy, pandas, openpyxl.

Run: python generate_cloudcare_raw_data.py [--output-dir PATH]
No personal information, latent variables, features or labels are exported.
Existing workbooks are never modified. Partial sets must match the deterministic
generation before missing workbooks can be added, preventing mixed FK universes.
"""
from pathlib import Path
import argparse
import hashlib
import json
import calendar
from datetime import datetime, timedelta
import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill

OUTPUT_DIR = Path(__file__).resolve().parents[1] / 'data' / 'raw' / 'cloudcare_20260831_seed42'
REFERENCE_DATE = datetime(2026, 8, 31)
START = datetime(2024, 1, 1)
END = REFERENCE_DATE.replace(hour=23, minute=59, second=59)
SEED = 42
SCHEMAS = {
 'user': 'user_id signup_date age_group region signup_channel status created_at updated_at',
 'plan': 'plan_id plan_name storage_limit_gb monthly_price billing_cycle is_paid is_active',
 'subscription': 'subscription_id user_id plan_id start_date next_billing_date end_date auto_renewal subscription_status created_at updated_at',
 'storage_usage_monthly': 'usage_id user_id usage_month storage_used_gb file_count photo_count video_count upload_size_gb download_size_gb created_at',
 'user_activity_daily': 'activity_id user_id activity_date login_count upload_count download_count share_count preview_count active_minutes',
 'device': 'device_id user_id device_type os_type sync_enabled last_sync_at registered_at',
 'payment_history': 'payment_id subscription_id payment_date amount payment_status payment_method retry_count',
 'support_ticket': 'ticket_id user_id category priority status created_at resolved_at reopened',
 'subscription_event': 'event_id user_id subscription_id event_type old_plan_id new_plan_id event_date',
}

def frame(name, rows):
    return pd.DataFrame(rows, columns=SCHEMAS[name].split())

def add_months(d, n):
    m = d.month - 1 + n
    y, m = d.year + m // 12, m % 12 + 1
    return d.replace(year=y, month=m, day=min(d.day, calendar.monthrange(y, m)[1]))

def create_output_directory(path):
    path.mkdir(parents=True, exist_ok=True)
    return path

def generate_plan():
    return frame('plan', [(i, name, limit, price, cycle, i != 1, True) for i, name, limit, price, cycle in [
        (1, 'FREE', 15, 0, 'MONTHLY'), (2, '100GB_M', 100, 2900, 'MONTHLY'),
        (3, '100GB_Y', 100, 2700, 'YEARLY'), (4, '200GB_M', 200, 4900, 'MONTHLY'),
        (5, '200GB_Y', 200, 4500, 'YEARLY'), (6, '2TB_M', 2000, 11900, 'MONTHLY'),
        (7, '2TB_Y', 2000, 10900, 'YEARLY')]])

def generate_hidden_profiles(rng):
    # All six requested drivers remain in memory only; no profile artifact is saved.
    return {100001+i: dict(zip(['engagement_score', 'satisfaction_score', 'price_sensitivity',
        'technical_friction', 'financial_risk', 'usage_growth'], rng.beta([2,4,2,2,1.4,2], [2,2,2,5,8,2])))
        | {'declining': bool(rng.random() < .24)} for i in range(5000)}

def generate_users(rng):
    regions = ['서울','경기','부산','인천','대구','대전','광주','울산','세종','강원','충북','충남','전북','전남','경북','경남','제주']
    weights = np.array([19,27,7,6,5,3,3,2,1,3,3,4,3,3,5,7,1], dtype=float)
    rows=[]
    for uid in range(100001,105001):
        signup=START+timedelta(days=int(rng.beta(1.07,1)*(REFERENCE_DATE-START).days))
        rows.append([uid,signup,rng.choice(['10대','20대','30대','40대','50대','60대 이상'],p=[.05,.25,.3,.2,.13,.07]),
            rng.choice(regions,p=weights/weights.sum()),rng.choice(['WEB','IOS','ANDROID'],p=[.25,.35,.4]),'ACTIVE',signup,signup])
    return frame('user',rows)

def generate_subscriptions(users, profiles, plans, rng):
    rows=[]
    eligible=users.loc[users.signup_date<=REFERENCE_DATE-timedelta(days=90),'user_id'].to_numpy()
    multiple=rng.choice(eligible,900,replace=False)
    counts={int(uid):2 if i<750 else 3 for i,uid in enumerate(multiple)}
    for u in users.itertuples(index=False):
        p=profiles[u.user_id]; span=(REFERENCE_DATE-u.signup_date).days
        count=counts.get(u.user_id,1)
        offsets=sorted(rng.choice(np.arange(30,span-14),size=count-1,replace=False).tolist()) if count>1 else []
        starts=[u.signup_date]+[u.signup_date+timedelta(days=int(x)) for x in offsets]
        weights=np.array([.28,.22,.11,.15,.08,.10,.06])
        pid=int(rng.choice(range(1,8),p=weights))
        for j,start in enumerate(starts):
            if j:
                old_limit=plans.loc[pid,'storage_limit_gb']
                down=rng.random()<(.25+.45*p['price_sensitivity'])
                choices=[k for k in range(1,8) if (plans.loc[k,'storage_limit_gb']<old_limit if down else plans.loc[k,'storage_limit_gb']>old_limit)]
                if not choices: choices=[k for k in range(1,8) if plans.loc[k,'storage_limit_gb']!=old_limit]
                pid=int(rng.choice(choices))
            if j<count-1:
                end=starts[j+1]-timedelta(days=1); status='EXPIRED' if rng.random()<.42 else 'CANCELLED'
            elif rng.random()<(.045+.055*p['price_sensitivity']+.04*(1-p['satisfaction_score'])+.025*p['financial_risk']):
                end=start+timedelta(days=int(rng.integers(0,max(1,(REFERENCE_DATE-start).days+1))))
                status='CANCELLED' if rng.random()<.77 else 'EXPIRED'
            else: end=None; status='ACTIVE'
            # Last stored renewal preference; ended status overrides execution.
            renewal=bool(rng.random()<(.89-.14*p['price_sensitivity']))
            cycle=12 if plans.loc[pid,'billing_cycle']=='YEARLY' else 1
            next_bill=None
            if status=='ACTIVE' and pid!=1:
                n=1; next_bill=add_months(start,cycle*n)
                while next_bill<=REFERENCE_DATE:
                    n+=1; next_bill=add_months(start,cycle*n)
            rows.append([len(rows)+1,u.user_id,pid,start,next_bill,end,renewal,status,start,end or REFERENCE_DATE])
    return frame('subscription',rows)

def plan_at(history, date):
    periods=history.attrs.get('periods')
    if periods is None:
        periods=[(r.start_date,r.end_date,int(r.plan_id)) for r in history.itertuples(index=False)]
        history.attrs['periods']=periods
    for start,end,pid in reversed(periods):
        if start<=date and (pd.isna(end) or end>=date): return pid
    return 1  # ended account retains free-tier access

def generate_storage_usage(users, histories, profiles, plans, rng):
    rows=[]; ratios={}
    for u in users.itertuples(index=False):
        p=profiles[u.user_id]; h=histories[u.user_id]
        used=float(plans.loc[int(h.iloc[0].plan_id),'storage_limit_gb'])*rng.uniform(.06,.60)
        month=u.signup_date.replace(day=1)
        while month<=REFERENCE_DATE:
            snapshot=min(add_months(month,1)-timedelta(days=1),REFERENCE_DATE)
            cap=float(plans.loc[plan_at(h,snapshot),'storage_limit_gb'])
            decline=.82 if p['declining'] and snapshot>=datetime(2026,6,1) else 1
            used=float(np.clip(used*(1+.07*(p['usage_growth']-.35)+rng.normal(0,.025))*decline+cap*.006*p['engagement_score'],0,cap*.98))
            files=int(used*rng.uniform(180,350)); photos=int(files*rng.uniform(.35,.65)); videos=int((files-photos)*rng.uniform(.03,.20))
            rows.append([len(rows)+1,u.user_id,month,round(used,2),files,photos,videos,
                round(cap*.07*p['engagement_score']*rng.lognormal(0,.4)*decline,2),round(cap*.045*p['engagement_score']*rng.lognormal(0,.5),2),snapshot.replace(hour=23)])
            ratios[(u.user_id,month)]=used/cap
            month=add_months(month,1)
    return frame('storage_usage_monthly',rows),ratios

def generate_user_activity(users, histories, profiles, rng):
    rows=[]
    for u in users.itertuples(index=False):
        p=profiles[u.user_id]; n=(REFERENCE_DATE-u.signup_date).days+1
        dates=np.arange(n); probability=np.full(n,.045+.22*p['engagement_score'])
        recent=dates>=max(0,n-90)
        if p['declining']: probability[recent]*=.27
        h=histories[u.user_id]; last=h.iloc[-1]
        if last.subscription_status!='ACTIVE': probability[dates>(last.end_date-u.signup_date).days]*=.18
        for offset in dates[rng.random(n)<probability]:
            date=u.signup_date+timedelta(days=int(offset)); scale=(.35 if p['declining'] and date>=datetime(2026,6,3) else 1)
            e=p['engagement_score']*scale
            login=1+int(rng.poisson(2.7*e)); upload=int(rng.poisson(8*e)) if rng.random()<.62 else 0
            download=int(rng.poisson(5*e)); share=int(rng.poisson(2*e)) if rng.random()<.18 else 0
            preview=int(rng.poisson(14*e)); minutes=int(np.clip(rng.gamma(2,9+24*e)+2*(upload+download),5,280))
            rows.append([len(rows)+1,u.user_id,date,login,upload,download,share,preview,minutes])
    return frame('user_activity_daily',rows)

def generate_devices(users, profiles, rng):
    rows=[]
    for u in users.itertuples(index=False):
        p=profiles[u.user_id]
        for _ in range(int(rng.choice([1,2,3],p=[.35,.48,.17]))):
            kind=rng.choice(['MOBILE','PC','TABLET'],p=[.55,.35,.10])
            os=rng.choice(['WINDOWS','MACOS'],p=[.75,.25]) if kind=='PC' else rng.choice(['IOS','ANDROID'],p=[.47,.53])
            registered=u.signup_date+timedelta(days=int(rng.integers(0,(REFERENCE_DATE-u.signup_date).days+1)))
            enabled=bool(rng.random()<.96-.5*p['technical_friction'])
            lag=int(rng.exponential(3+90*p['technical_friction']+(0 if enabled else 45)))
            sync=max(registered,REFERENCE_DATE-timedelta(days=lag)) if rng.random()>.05+.12*p['technical_friction'] else None
            rows.append([len(rows)+1,u.user_id,kind,os,enabled,sync,registered])
    return frame('device',rows)

def generate_payment_history(subs, users, profiles, plans, rng):
    rows=[]; channels=users.set_index('user_id').signup_channel
    for s in subs.itertuples(index=False):
        if s.plan_id==1: continue
        plan=plans.loc[s.plan_id]; cycle=12 if plan.billing_cycle=='YEARLY' else 1
        amount=float(plan.monthly_price)*cycle; finish=s.end_date if pd.notna(s.end_date) else REFERENCE_DATE
        n=0; date=s.start_date
        while date<=finish:
            risk=profiles[s.user_id]['financial_risk']; failed=.018+.17*risk
            status=rng.choice(['SUCCESS','FAILED','REFUNDED'],p=[1-failed-.014,failed,.014])
            primary={'WEB':'CARD','IOS':'APP_STORE','ANDROID':'PLAY_STORE'}[channels[s.user_id]]
            method=primary if rng.random()<.90 else rng.choice(['CARD','APP_STORE','PLAY_STORE'])
            retry=int(rng.choice([0,1,2,3],p=[.10,.55,.25,.10])) if status=='FAILED' else int(rng.random()<.035)
            rows.append([len(rows)+1,s.subscription_id,date.replace(hour=12),amount,status,method,retry])
            n+=1; date=add_months(s.start_date,cycle*n)
    return frame('payment_history',rows)

def generate_support_tickets(users, profiles, ratios, rng):
    rows=[]
    for u in users.itertuples(index=False):
        p=profiles[u.user_id]; age=(REFERENCE_DATE-u.signup_date).days
        probability=(.16+.39*p['technical_friction']+.16*(1-p['satisfaction_score']))*min(1,.25+age/180)
        if rng.random()>probability: continue
        count=1+int(rng.poisson(2.2+4*p['technical_friction']+2*(1-p['satisfaction_score'])))
        if rng.random()<.018: count+=int(rng.integers(25,65))
        for _ in range(count):
            date=u.signup_date+timedelta(days=int(rng.integers(0,age+1)),hours=int(rng.integers(0,20)))
            ratio=ratios[(u.user_id,date.replace(day=1,hour=0))]
            w=np.array([.25+2*p['financial_risk'],.25+2*p['technical_friction'],.15+1.5*ratio,.25+1.5*p['technical_friction']])
            category=rng.choice(['결제','동기화','저장공간','오류'],p=w/w.sum())
            priority=rng.choice(['LOW','NORMAL','HIGH','URGENT'],p=[.15,.60,.20,.05])
            resolved=date+timedelta(hours=float(rng.gamma(2,12+35*p['technical_friction'])))
            if resolved>END or rng.random()<.025: resolved=None
            reopened=bool(rng.random()<.025+.25*p['technical_friction']+.08*(1-p['satisfaction_score']))
            rows.append([len(rows)+1,u.user_id,category,priority,'RESOLVED' if resolved else 'OPEN',date,resolved,reopened])
    return frame('support_ticket',rows)

def generate_subscription_events(histories, plans):
    rows=[]
    def event(s,kind,old,new,date): rows.append([len(rows)+1,s.user_id,s.subscription_id,kind,old,new,date])
    for history in histories.values():
        previous=None
        for s in history.itertuples(index=False):
            if previous is not None:
                kind='UPGRADE' if plans.loc[s.plan_id,'storage_limit_gb']>plans.loc[previous.plan_id,'storage_limit_gb'] else 'DOWNGRADE'
                event(s,kind,previous.plan_id,s.plan_id,s.start_date)
            finish=s.end_date if pd.notna(s.end_date) else REFERENCE_DATE
            cycle=12 if plans.loc[s.plan_id,'billing_cycle']=='YEARLY' else 1
            # Administrative lifecycle period renewal, independent of payment outcome.
            # Free access has no billed contract renewal. Monthly rolling contracts
            # receive lifecycle renewal every three months, yearly every twelve.
            step=12 if cycle==12 else 3
            if s.plan_id!=1:
                n=step; date=add_months(s.start_date,n)
                while date<=finish:
                    event(s,'RENEW',s.plan_id,s.plan_id,date)
                    n+=step; date=add_months(s.start_date,n)
            if s.subscription_status=='CANCELLED': event(s,'CANCEL',s.plan_id,None,s.end_date)
            previous=s
    return frame('subscription_event',rows)

def validate_primary_keys(tables):
    for name, df in tables.items():
        pk=df.iloc[:,0]; assert pk.notna().all() and pk.is_unique, name+' PK'

def validate_foreign_keys(t):
    targets={'user_id':set(t['user'].user_id),'plan_id':set(t['plan'].plan_id),
             'subscription_id':set(t['subscription'].subscription_id)}
    for name,df in t.items():
        for col in df.columns[1:]:
            target='plan_id' if col in ['old_plan_id','new_plan_id'] else col
            if target in targets:
                assert set(df[col].dropna())<=targets[target], (name,col)
                if col not in ['old_plan_id','new_plan_id']: assert df[col].notna().all()
    e=t['subscription_event'].merge(t['subscription'][['subscription_id','user_id']],on='subscription_id',suffixes=('','_sub'))
    assert (e.user_id==e.user_id_sub).all()

def validate_business_rules(t, clean=False):
    u=t['user'].set_index('user_id'); s=t['subscription']; plans=t['plan'].set_index('plan_id')
    assert t['plan'].equals(generate_plan())
    assert (s.start_date>=s.user_id.map(u.signup_date)).all()
    assert (s.loc[s.end_date.notna(),'end_date']>=s.loc[s.end_date.notna(),'start_date']).all()
    assert s.loc[s.subscription_status=='ACTIVE','end_date'].isna().all()
    assert s.loc[s.subscription_status!='ACTIVE','end_date'].notna().all()
    for _,h in s.groupby('user_id'):
        h=h.sort_values('start_date'); a=h.iloc[:-1]; b=h.iloc[1:]
        assert (a.end_date.to_numpy()<b.start_date.to_numpy()).all()
    for name,df in t.items():
        for c in df:
            if pd.api.types.is_datetime64_any_dtype(df[c]):
                vals=df[c].dropna(); assert (vals>=START).all(), (name,c)
                if c!='next_billing_date': assert (vals<=END).all(), (name,c)
    d=t['device']; valid=d.last_sync_at.notna()
    assert (d.loc[valid,'last_sync_at']>=d.loc[valid,'registered_at']).all()
    assert (d.registered_at>=d.user_id.map(u.signup_date)).all()
    for kind,allowed in {'MOBILE':['IOS','ANDROID'],'TABLET':['IOS','ANDROID'],'PC':['WINDOWS','MACOS']}.items():
        assert d.loc[(d.device_type==kind)&d.os_type.notna(),'os_type'].isin(allowed).all()
    tickets=t['support_ticket']; resolved=tickets.status=='RESOLVED'
    assert tickets.loc[~resolved,'resolved_at'].isna().all()
    assert tickets.loc[resolved,'resolved_at'].notna().all()
    assert (tickets.loc[resolved,'resolved_at']>=tickets.loc[resolved,'created_at']).all()
    assert (tickets.created_at>=tickets.user_id.map(u.signup_date)).all()
    a=t['user_activity_daily']; assert (a.activity_date>=a.user_id.map(u.signup_date)).all()
    usage=t['storage_usage_monthly']; counts=usage[['file_count','photo_count','video_count']].dropna()
    assert (counts.photo_count+counts.video_count<=counts.file_count).all()
    assert (usage.created_at>=usage.user_id.map(u.signup_date)).all()
    p=t['payment_history'].merge(s,on='subscription_id'); assert (p.plan_id!=1).all()
    assert (p.payment_date>=p.start_date).all()
    assert (p.payment_date.dt.normalize()<=p.end_date.fillna(REFERENCE_DATE)).all()
    if clean:
        expected=p.plan_id.map(plans.monthly_price)*p.plan_id.map(plans.billing_cycle).map({'MONTHLY':1,'YEARLY':12})
        assert (p.amount==expected).all()
        histories={uid:h for uid,h in s.groupby('user_id')}
        for row in usage.itertuples(index=False):
            assert 0<=row.storage_used_gb<=plans.loc[plan_at(histories[row.user_id],row.created_at.normalize()),'storage_limit_gb']*.98+.01
    e=t['subscription_event'].merge(s,on=['subscription_id','user_id'])
    for kind,op in [('UPGRADE',np.greater),('DOWNGRADE',np.less)]:
        v=e[e.event_type==kind]
        assert op(v.new_plan_id.map(plans.storage_limit_gb),v.old_plan_id.map(plans.storage_limit_gb)).all()
        assert (v.new_plan_id==v.plan_id).all() and (v.event_date==v.start_date).all()
    c=e[e.event_type=='CANCEL']; assert (c.subscription_status=='CANCELLED').all()
    assert c.new_plan_id.isna().all() and (c.event_date==c.end_date).all()
    assert set(c.subscription_id)==set(s.loc[s.subscription_status=='CANCELLED','subscription_id'])
    r=e[e.event_type=='RENEW']; assert (r.old_plan_id==r.new_plan_id).all()
    assert (e.event_date>=e.start_date).all() and (e.event_date<=e.end_date.fillna(REFERENCE_DATE)).all()

def inject_missing_values(t,rng,audit):
    columns={'user':{'region':.015,'age_group':.008,'signup_channel':.004},
        'storage_usage_monthly':{c:.005 for c in ['storage_used_gb','file_count','photo_count','video_count','upload_size_gb','download_size_gb']},
        'user_activity_daily':{c:.005 for c in SCHEMAS['user_activity_daily'].split()[3:]},
        'device':{'os_type':.008},'payment_history':{'payment_method':.005},'support_ticket':{'category':.005}}
    for name,cols in columns.items():
        for col,rate in cols.items():
            df=t[name]; idx=rng.choice(df.index,size=round(len(df)*rate),replace=False)
            if pd.api.types.is_integer_dtype(df[col]): df[col]=df[col].astype('Int64')
            df.loc[idx,col]=pd.NA
            audit[name]['inserted_nulls']+=len(idx)

def inject_outliers(t,rng,audit):
    for name,col,rate,low,high in [('user_activity_daily','active_minutes',.012,300,801),
        ('storage_usage_monthly','download_size_gb',.015,3,9),('payment_history','amount',.002,1.1,1.25)]:
        df=t[name]; eligible=df.index[df[col].notna()]; idx=rng.choice(eligible,size=round(len(df)*rate),replace=False)
        if name=='user_activity_daily': df.loc[idx,col]=rng.integers(low,high,len(idx))
        else: df.loc[idx,col]=(df.loc[idx,col].astype(float)*rng.uniform(low,high,len(idx))).round(2)
        audit[name]['outliers_or_errors']+=len(idx)
    # Very rare telemetry counter error; protected count/date constraints stay intact.
    df=t['user_activity_daily']; idx=rng.choice(df.index[df.login_count.notna()],round(len(df)*.001),replace=False)
    df.loc[idx,'login_count']=-1; audit['user_activity_daily']['outliers_or_errors']+=len(idx)

def inject_logical_duplicates(t,rng,audit):
    for name,rate in [('storage_usage_monthly',.007),('user_activity_daily',.005),('payment_history',.005),('support_ticket',.004)]:
        df=t[name]; n=round(len(df)*rate); dup=df.loc[rng.choice(df.index,n,replace=False)].copy()
        dup.iloc[:,0]=np.arange(int(df.iloc[:,0].max())+1,int(df.iloc[:,0].max())+n+1)
        t[name]=pd.concat([df,dup],ignore_index=True); audit[name]['logical_duplicates']=n

def normalized_rows(df):
    for row in df.itertuples(index=False,name=None):
        result=[]
        for value in row:
            if pd.isna(value): value=None
            elif isinstance(value,pd.Timestamp): value=value.to_pydatetime().replace(microsecond=(value.microsecond//1000)*1000)
            elif isinstance(value,np.generic): value=value.item()
            result.append(value)
        yield tuple(result)

def verify_saved_file(path,name,df):
    wb=load_workbook(path,read_only=True,data_only=True)
    try:
        assert wb.sheetnames==[name], f'{path}: sheet mismatch'
        rows=wb[name].iter_rows(min_col=1,max_col=len(df.columns),values_only=True)
        assert next(rows)==tuple(df.columns), f'{path}: columns mismatch'
        count=0
        for expected in normalized_rows(df):
            actual=next(rows,None)
            assert actual is not None and len(actual)==len(expected), f'{path}: missing row'
            for x,y in zip(actual,expected):
                if isinstance(y,float): assert x is not None and abs(x-y)<1e-8, (path,count)
                else: assert x==y, (path,count,x,y)
            count+=1
        assert next(rows,None) is None, f'{path}: excess rows'
        return count
    finally: wb.close()

def save_excel_if_not_exists(path,name,df):
    if path.exists():
        print(f'[SKIP] {path.name} - already exists',flush=True); return 'SKIP'
    # xb uses O_EXCL: even a concurrent creator cannot be overwritten.
    try: handle=path.open('xb')
    except FileExistsError:
        print(f'[SKIP] {path.name} - already exists',flush=True); return 'SKIP'
    with handle:
        wb=Workbook(write_only=True); ws=wb.create_sheet(name); ws.freeze_panes='A2'
        from openpyxl.utils import get_column_letter
        for i,col in enumerate(df.columns,1):
            ws.column_dimensions[get_column_letter(i)].width=max(23 if pd.api.types.is_datetime64_any_dtype(df[col]) else 17,len(col)+3)
        headers=[]
        for col in df:
            cell=WriteOnlyCell(ws,col); cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='24476A'); headers.append(cell)
        ws.append(headers)
        for row in normalized_rows(df):
            cells=[]
            for col,value in zip(df.columns,row):
                cell=WriteOnlyCell(ws,value)
                if isinstance(value,datetime): cell.number_format='yyyy-mm-dd hh:mm:ss' if col.endswith('_at') or col in ['payment_date','event_date'] else 'yyyy-mm-dd'
                elif col.endswith('_gb') or col in ['amount','monthly_price']: cell.number_format='0.00'
                cells.append(cell)
            ws.append(cells)
        ws.auto_filter.ref=f'A1:{get_column_letter(len(df.columns))}{len(df)+1}'
        wb.save(handle)
    print(f'[CREATED] {path.name} - {len(df):,} rows',flush=True)
    return 'CREATED'

def main():
    parser=argparse.ArgumentParser(); parser.add_argument('--output-dir',type=Path,default=OUTPUT_DIR)
    args=parser.parse_args(); output=create_output_directory(args.output_dir.resolve()); rng=np.random.default_rng(SEED)
    print('===== CloudCare Raw Dataset Generation =====',flush=True)
    profiles=generate_hidden_profiles(rng); t={'user':generate_users(rng),'plan':generate_plan()}
    plans=t['plan'].set_index('plan_id')
    t['subscription']=generate_subscriptions(t['user'],profiles,plans,rng)
    histories={uid:h.sort_values('start_date') for uid,h in t['subscription'].groupby('user_id')}
    t['storage_usage_monthly'],ratios=generate_storage_usage(t['user'],histories,profiles,plans,rng)
    print('[GENERATED] subscriptions and continuous monthly storage',flush=True)
    t['user_activity_daily']=generate_user_activity(t['user'],histories,profiles,rng)
    t['device']=generate_devices(t['user'],profiles,rng)
    t['payment_history']=generate_payment_history(t['subscription'],t['user'],profiles,plans,rng)
    t['support_ticket']=generate_support_tickets(t['user'],profiles,ratios,rng)
    t['subscription_event']=generate_subscription_events(histories,plans)
    for name,df in t.items(): print(f'[GENERATED] {name}: {len(df):,} clean rows',flush=True)
    recent=set(t['user_activity_daily'].loc[t['user_activity_daily'].activity_date>=REFERENCE_DATE-timedelta(days=30),'user_id'])
    active=set(t['subscription'].loc[t['subscription'].subscription_status=='ACTIVE','user_id'])
    t['user']['status']=['ACTIVE' if uid in recent or uid in active else 'INACTIVE' for uid in t['user'].user_id]
    t['user']['updated_at']=REFERENCE_DATE
    print('[VALIDATING] clean PK/FK/date/business rules',flush=True)
    validate_primary_keys(t); validate_foreign_keys(t); validate_business_rules(t,clean=True)
    audit={n:dict(inserted_nulls=0,logical_duplicates=0,missing_logs=0,outliers_or_errors=0) for n in t}
    for n in ['storage_usage_monthly','user_activity_daily']:
        df=t[n]; count=round(len(df)*.02); t[n]=df.drop(rng.choice(df.index,count,replace=False)).reset_index(drop=True); audit[n]['missing_logs']=count
    inject_missing_values(t,rng,audit); inject_outliers(t,rng,audit); inject_logical_duplicates(t,rng,audit)
    validate_primary_keys(t); validate_foreign_keys(t); validate_business_rules(t)
    for n,lo,hi in [('subscription',6000,8000),('storage_usage_monthly',30000,100000),('user_activity_daily',30000,1048575),('device',7000,12000),('payment_history',10000,60000),('support_ticket',5000,10000),('subscription_event',5000,30000)]:
        assert lo<=len(t[n])<=hi, (n,len(t[n]))
    paths={n:output/f'5.{i}_{n}.xlsx' for i,n in enumerate(SCHEMAS,1)}
    # Read-only compatibility check before creating any missing member of a set.
    for n,path in paths.items():
        if path.exists():
            print(f'[CHECK EXISTING] {path.name}',flush=True)
            verify_saved_file(path,n,t[n])
    results={}
    for n,path in paths.items():
        status=save_excel_if_not_exists(path,n,t[n]); verified=verify_saved_file(path,n,t[n])
        results[n]={'path':str(path),'rows':verified,'missing_cells':int(t[n].isna().sum().sum()),
            'missing_by_column':{c:int(v) for c,v in t[n].isna().sum().items()},**audit[n],'file_status':status,
            'bytes':path.stat().st_size,'sha256':hashlib.sha256(path.read_bytes()).hexdigest()}
    report={'seed':SEED,'reference_date':'2026-08-31','validation':{'clean':'PASS','raw':'PASS','saved_excel_roundtrip':'PASS'},'tables':results}
    print(json.dumps(report,ensure_ascii=False,indent=2),flush=True)
    report_path=output/'generation_report.json'
    if not report_path.exists():
        try:
            with report_path.open('x',encoding='utf-8') as f: json.dump(report,f,ensure_ascii=False,indent=2)
        except FileExistsError: pass
    print('[PASS] PK uniqueness / FK integrity / Date consistency / Business rules / Excel roundtrip',flush=True)

if __name__=='__main__': main()
