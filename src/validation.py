"""Fail before writing on structural errors; preserve documented raw anomalies."""
from decimal import Decimal, InvalidOperation
import datetime as dt
import math
import pandas as pd
from openpyxl import load_workbook
from sqlalchemy import inspect
from src.schema import ROOT, SPECS, ORDER, metadata

def fail_rows(name, col, series, mask, reason):
    examples = [{'excel_row': int(i)+2, 'value': str(series.loc[i])} for i in series.index[mask][:5]]
    raise ValueError(f'{name}.{col}: {reason}; count={int(mask.sum())}; examples={examples}')

def validate_frame(name, frame):
    spec = SPECS[name]
    expected = list(spec['columns'])
    if list(frame.columns) != expected:
        raise ValueError(f'{name}: expected columns={expected}, actual={list(frame.columns)}')
    df = frame.copy()
    for col, info in spec['columns'].items():
        series = df[col]
        missing = series.isna()
        if not info['nullable'] and missing.any():
            fail_rows(name, col, series, missing, 'required NULL')
        kind = info['type']
        if kind == 'DATE' or kind.startswith('DATETIME'):
            # Numeric Excel serials are not guessed as nanosecond timestamps.
            invalid_type = series.map(lambda x: not pd.isna(x) and not isinstance(x, (str, dt.date, dt.datetime, pd.Timestamp)))
            parsed = pd.to_datetime(series.where(~invalid_type), errors='coerce', format='mixed')
            invalid = (~missing & parsed.isna()) | invalid_type
            if invalid.any():
                fail_rows(name, col, series, invalid, 'invalid date')
            if kind == 'DATE' and ((parsed != parsed.dt.normalize()) & ~missing).any():
                fail_rows(name, col, series, (parsed != parsed.dt.normalize()) & ~missing, 'DATE would lose time')
            invalid = ~missing & ((parsed.dt.year < 1000) | (parsed.dt.nanosecond != 0))
            if invalid.any():
                fail_rows(name, col, series, invalid, 'date range or fractional seconds')
            df[col] = parsed
        elif kind in {'INT','BIGINT','DOUBLE','BOOLEAN'} or kind.startswith('DECIMAL'):
            numeric = pd.to_numeric(series, errors='coerce')
            invalid = ~missing & (numeric.isna() | ~numeric.map(lambda x: pd.isna(x) or math.isfinite(x)))
            if kind in {'INT','BIGINT','BOOLEAN'}:
                low, high = (-2**31, 2**31-1) if kind == 'INT' else (-2**63, 2**63-1)
                invalid |= ~missing & ((numeric % 1 != 0) | (numeric < low) | (numeric > high))
                if kind == 'BOOLEAN':
                    invalid |= ~missing & ~numeric.isin([0,1])
            if invalid.any():
                fail_rows(name, col, series, invalid, 'invalid numeric value')
            if kind.startswith('DECIMAL'):
                precision, scale = map(int, kind.split('(')[1][:-1].split(','))
                converted = []
                for i, value in series.items():
                    if pd.isna(value):
                        converted.append(None)
                        continue
                    number = Decimal(str(value))
                    try:
                        valid = number.is_finite() and abs(number) < Decimal(10)**(precision-scale) and number == number.quantize(Decimal(10)**(-scale))
                    except InvalidOperation:
                        valid = False
                    if not valid:
                        raise ValueError(f'{name}.{col}: Excel row {i+2}, invalid {kind}: {value}')
                    converted.append(number)
                df[col] = converted
            else:
                df[col] = numeric
        else:
            limit = int(kind.split('(')[1][:-1])
            invalid = series.map(lambda x: not pd.isna(x) and (not isinstance(x, str) or len(x) > limit))
            if invalid.any():
                fail_rows(name, col, series, invalid, 'invalid string or length')
    pk = spec['pk']
    if df[pk].duplicated(keep=False).any():
        fail_rows(name, pk, df[pk], df[pk].duplicated(keep=False), 'duplicate PK')
    if name == 'plan' and df['plan_name'].str.casefold().duplicated(keep=False).any():
        fail_rows(name, 'plan_name', df['plan_name'], df['plan_name'].str.casefold().duplicated(keep=False), 'duplicate plan name')
    return df

def read_sources():
    frames = {}
    for name in ORDER:
        spec = SPECS[name]
        path = ROOT/spec['file']
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            if book.sheetnames != [spec['sheet']]:
                raise ValueError(f'{name}: unexpected sheets {book.sheetnames}')
            header = next(book[spec['sheet']].iter_rows(values_only=True))
            if list(header) != list(spec['columns']):
                raise ValueError(f'{name}: unexpected Excel header {header}')
        finally:
            book.close()
        df = pd.read_excel(path, sheet_name=spec['sheet'], engine='openpyxl', keep_default_na=False, na_values=[''])
        frames[name] = validate_frame(name, df)
        print(f'[VALID] {name}: {len(df):,} rows', flush=True)
    for name, spec in SPECS.items():
        for col, (parent, pk) in spec['fk'].items():
            series = frames[name][col]
            invalid = series.notna() & ~series.isin(frames[parent][pk])
            if invalid.any():
                fail_rows(name, col, series, invalid, f'FK missing from {parent}.{pk}')
    return frames

def validate_db_schema(connection):
    inspector = inspect(connection)
    for name, spec in SPECS.items():
        if not inspector.has_table(name):
            raise ValueError(f'{name}: table missing; initialize db/init/01_schema.sql first')
        columns = inspector.get_columns(name)
        if [col['name'] for col in columns] != list(spec['columns']):
            raise ValueError(f'{name}: Excel/DB columns mismatch')
        for col in columns:
            expected = metadata.tables[name].c[col['name']]
            actual = col['type']
            want = expected.type
            # MySQL reflects BOOLEAN as TINYINT(1), DOUBLE as DOUBLE.
            from sqlalchemy import Boolean, Float, Numeric, String
            if isinstance(want, Boolean):
                match = actual.__class__.__name__ == 'TINYINT' and getattr(actual, 'display_width', None) == 1
            elif isinstance(want, Float):
                match = actual.__class__.__name__ == 'DOUBLE'
            else:
                match = actual._type_affinity is want._type_affinity
                if spec['columns'][col['name']]['type'] in {'INT', 'BIGINT'}:
                    match &= actual.__class__.__name__ == ('BIGINT' if spec['columns'][col['name']]['type'] == 'BIGINT' else 'INTEGER')
                if spec['columns'][col['name']]['type'] == 'DATETIME(6)':
                    match &= getattr(actual, 'fsp', None) == 6
                if isinstance(want, String):
                    match &= actual.length == want.length
                elif isinstance(want, Numeric):
                    match &= actual.precision == want.precision and actual.scale == want.scale
            if getattr(actual, 'unsigned', False):
                match = False
            if not match or col['nullable'] != expected.nullable:
                raise ValueError(f'{name}.{col["name"]}: DB type/nullability mismatch: {actual}, nullable={col["nullable"]}; expected {want}, nullable={expected.nullable}')
        if inspector.get_pk_constraint(name)['constrained_columns'] != [spec['pk']]:
            raise ValueError(f'{name}: DB PK mismatch')
        actual_fks = {(tuple(f['constrained_columns']),f['referred_table'],tuple(f['referred_columns'])) for f in inspector.get_foreign_keys(name)}
        expected_fks = {((col,), parent, (pk,)) for col,(parent,pk) in spec['fk'].items()}
        if actual_fks != expected_fks:
            raise ValueError(f'{name}: DB FK mismatch')
        if name == 'plan' and not any(item['column_names'] == ['plan_name'] for item in inspector.get_unique_constraints(name)):
            raise ValueError('plan: unique plan_name constraint missing')
